"""Extract bounded, readable text from imported office documents and email."""

import io
from email import policy
from email.parser import BytesParser
from pathlib import Path


SUPPORTED = {".pdf", ".doc", ".docx", ".xls", ".xlsx", ".ppt", ".pptx", ".eml",
             ".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}
MAX_TEXT = 100_000


def extract(path: Path, *, check_active=lambda: None, progress=None) -> tuple[str, list[tuple[str, bytes]]]:
    check_active()
    suffix = path.suffix.lower()
    if suffix in {".doc", ".ppt"}:
        import legacy_office
        return legacy_office.extract(path), []
    if suffix in {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}:
        import local_ocr
        return local_ocr.file_text(path, check_active, progress)[:MAX_TEXT], []
    if suffix == ".pptx":
        from pptx import Presentation
        parts = []
        count = 0
        def shape_text(shape):
            if shape.has_text_frame:
                yield shape.text
            if shape.has_table:
                for row in shape.table.rows:
                    yield "\t".join(cell.text for cell in row.cells)
            for child in getattr(shape, "shapes", []):
                yield from shape_text(child)
        slides = Presentation(str(path)).slides
        for index, slide in enumerate(slides):
            if index >= 200 or count >= MAX_TEXT:
                break
            check_active()
            values = [value for shape in slide.shapes for value in shape_text(shape)]
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    values.append("备注：" + notes)
            values = [value for value in values if value.strip()]
            if values:
                parts.append(f"第 {index + 1} 页")
            for value in values:
                parts.append(value[:MAX_TEXT - count])
                count += len(parts[-1]) + 1
                if count >= MAX_TEXT:
                    break
            if progress:
                progress(index + 1, min(200, len(slides)))
        return "\n".join(parts)[:MAX_TEXT], []
    if suffix == ".xls":
        import xlrd
        book = xlrd.open_workbook(str(path), on_demand=True)
        try:
            parts = []
            count = 0
            for sheet in book.sheets()[:50]:
                check_active()
                parts.append(f"工作表：{sheet.name}")
                for row in range(min(sheet.nrows, 10000)):
                    check_active()
                    text = "\t".join(str(sheet.cell_value(row, col)) for col in range(min(sheet.ncols, 200)))
                    if text.strip():
                        parts.append(text[:MAX_TEXT - count])
                        count += len(parts[-1]) + 1
                    if count >= MAX_TEXT:
                        break
                if count >= MAX_TEXT:
                    break
            return "\n".join(parts)[:MAX_TEXT], []
        finally:
            book.release_resources()
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        if reader.is_encrypted:
            raise ValueError("PDF 已加密，请先解锁后重新导入")
        pages = []
        character_count = 0
        has_text = False
        for index, page in enumerate(reader.pages[:200], 1):
            check_active()
            if character_count >= MAX_TEXT:
                break
            page_text = (page.extract_text(extraction_mode="layout") or "") if "/Contents" in page else ""
            if not page_text.strip() and "/Contents" in page:
                import local_ocr
                page_text = local_ocr.pdf_page_text(path, index - 1, check_active)
            has_text = has_text or bool(page_text.strip())
            section = f"第 {index} 页\n{page_text}"[:MAX_TEXT - character_count]
            pages.append(section)
            character_count += len(section) + 2
            if progress:
                progress(index, min(200, len(reader.pages)))
        text = "\n\n".join(pages).strip()
        if not has_text:
            raise ValueError("扫描版或空白 PDF 经 OCR 后未识别到文字，请检查扫描清晰度")
        return text, []
    if suffix == ".docx":
        from docx import Document

        document = Document(str(path))
        parts = []
        character_count = 0
        for paragraph in document.paragraphs:
            check_active()
            if character_count >= MAX_TEXT:
                break
            if paragraph.text.strip():
                value = paragraph.text[:MAX_TEXT - character_count]
                parts.append(value)
                character_count += len(value) + 1
        for table in document.tables:
            if character_count >= MAX_TEXT:
                break
            for row in table.rows:
                check_active()
                if character_count >= MAX_TEXT:
                    break
                value = "\t".join(cell.text for cell in row.cells)[:MAX_TEXT - character_count]
                parts.append(value)
                character_count += len(value) + 1
        return "\n".join(parts), []
    if suffix == ".xlsx":
        from openpyxl import load_workbook

        book = load_workbook(str(path), read_only=True, data_only=True)
        try:
            parts = []
            character_count = 0
            for sheet in book.worksheets[:50]:
                parts.append(f"工作表：{sheet.title}")
                character_count += len(parts[-1])
                for row in sheet.iter_rows(max_row=10000, max_col=200, values_only=True):
                    check_active()
                    if character_count >= MAX_TEXT:
                        break
                    line = "\t".join(str(value) if value is not None else "" for value in row)
                    if line.strip():
                        line = line[:MAX_TEXT - character_count]
                        parts.append(line)
                        character_count += len(line) + 1
                    if character_count >= MAX_TEXT:
                        break
                if character_count >= MAX_TEXT:
                    break
            return "\n".join(parts)[:MAX_TEXT], []
        finally:
            book.close()
    if suffix == ".eml":
        message = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
        headers = [f"{key}：{message.get(key, '')}" for key in ("From", "To", "Date", "Subject")]
        body = message.get_body(preferencelist=("plain", "html")) if message.is_multipart() else message
        content = body.get_content() if body else ""
        if body and body.get_content_type() == "text/html":
            from html.parser import HTMLParser

            class PlainText(HTMLParser):
                def __init__(self):
                    super().__init__()
                    self.parts = []
                    self.ignored = 0

                def handle_starttag(self, tag, attrs):
                    if tag in ("script", "style", "head"):
                        self.ignored += 1

                def handle_endtag(self, tag):
                    if tag in ("script", "style", "head"):
                        self.ignored = max(0, self.ignored - 1)

                def handle_data(self, data):
                    if not self.ignored:
                        self.parts.append(data)

            parser = PlainText()
            parser.feed(str(content))
            content = " ".join(parser.parts)
        attachments = []
        attachment_bytes = 0
        for part in message.iter_attachments():
            filename = part.get_filename()
            if filename:
                data = part.get_payload(decode=True)
                if data and len(attachments) < 20 and len(data) <= 25 * 1024 * 1024 and attachment_bytes + len(data) <= 100 * 1024 * 1024:
                    attachments.append((filename, data))
                    attachment_bytes += len(data)
        return ("\n".join(headers) + "\n\n" + str(content))[:MAX_TEXT], attachments
    raise ValueError("该文件类型暂不支持内容解析")
