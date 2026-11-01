"""Offline regression tests; no model API or application database is used."""

import shutil
import struct
import tempfile
import threading
import time
import unittest
import wave
from pathlib import Path
from unittest.mock import patch

import audio_pipeline
import document_ingest
import legacy_office


class AudioTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def fake_ffmpeg(self, args, **kwargs):
        manifest = Path(args[args.index('-segment_list') + 1])
        rows = []
        for index in range(5):
            path = manifest.parent / f'chunk-{index:06d}.mp3'
            path.write_bytes(str(index).encode())  # Include a tiny tail chunk.
            rows.append(f'{path.name},{index * 55},{index * 55 + (1 if index == 4 else 55)}')
        manifest.write_text('\n'.join(rows), encoding='utf-8')
        class Process:
            returncode = 0
            def poll(self):
                return 0
        return Process()

    def run_audio(self, recognize, check=lambda: None):
        return audio_pipeline.transcribe(self.root / 'source.wav', ffmpeg='ffmpeg',
            cache_dir=self.root / 'cache', cache_key='test-model', recognize=recognize, check_active=check)

    def test_parallel_bounded_order_and_retry_checkpoints(self):
        lock = threading.Lock()
        active = peak = 0
        called = []
        def recognize(data):
            nonlocal active, peak
            index = int(data)
            with lock:
                active += 1
                peak = max(peak, active)
                called.append(index)
            time.sleep((5 - index) * .015)
            with lock:
                active -= 1
            return f'text-{index}'
        with patch.object(audio_pipeline.subprocess, 'Popen', self.fake_ffmpeg):
            result = self.run_audio(recognize)
            self.assertEqual([row['text'] for row in result], [f'text-{i}' for i in range(5)])
            self.assertEqual(peak, 3)
            self.assertEqual(result[-1]['end'], 221)
            called.clear()
            self.assertEqual(self.run_audio(recognize), result)
            self.assertEqual(called, [])

    def test_transient_retry_and_cancel_does_not_cache_result(self):
        attempts = {}
        def recognize(data):
            index = int(data)
            attempts[index] = attempts.get(index, 0) + 1
            if index == 0 and attempts[index] == 1:
                raise RuntimeError('HTTP 429')
            return f'text-{index}'
        with patch.object(audio_pipeline.subprocess, 'Popen', self.fake_ffmpeg):
            self.run_audio(recognize)
            self.assertEqual(attempts[0], 2)
        cancelled = threading.Event()
        class Cancelled(Exception):
            pass
        def check():
            if cancelled.is_set():
                raise Cancelled()
        def cancel(data):
            cancelled.set()
            return 'should not save'
        for checkpoint in (self.root / 'cache').glob('*.json'):
            checkpoint.unlink()
        with patch.object(audio_pipeline.subprocess, 'Popen', self.fake_ffmpeg), self.assertRaises(Cancelled):
            self.run_audio(cancel, check)
        time.sleep(.05)
        self.assertEqual(list((self.root / 'cache').glob('*.json')), [])

    def test_failure_keeps_completed_chunks_for_retry(self):
        def fail(data):
            if data == b'4':
                time.sleep(.1)
                raise ValueError('bad response')
            return 'ok-' + data.decode()
        with patch.object(audio_pipeline.subprocess, 'Popen', self.fake_ffmpeg):
            with self.assertRaisesRegex(ValueError, 'bad response'):
                self.run_audio(fail)
            called = []
            def recover(data):
                called.append(data)
                return 'recovered'
            result = self.run_audio(recover)
            self.assertEqual(called, [b'4'])
            self.assertEqual(result[-1]['text'], 'recovered')

    def test_real_ffmpeg_offsets_and_short_tail(self):
        ffmpeg = shutil.which('ffmpeg') or str(Path(__file__).resolve().parents[1] / 'ffmpeg' / 'ffmpeg.exe')
        if not Path(ffmpeg).is_file() and not shutil.which(ffmpeg):
            self.skipTest('FFmpeg not installed')
        source = self.root / 'source.wav'
        with wave.open(str(source), 'wb') as output:
            output.setparams((1, 2, 16000, 0, 'NONE', 'not compressed'))
            output.writeframes(b'\0\0' * (16000 * 111))
        result = audio_pipeline.transcribe(source, ffmpeg=ffmpeg, cache_dir=self.root / 'cache',
            cache_key='real-segments', recognize=lambda data: 'recognized', check_active=lambda: None)
        self.assertEqual(len(result), 3)
        self.assertLess(result[-1]['end'] - result[-1]['start'], 2)
        self.assertTrue(all(0 < row['end'] - row['start'] < 60 for row in result))
        self.assertAlmostEqual(result[1]['start'], result[0]['end'], places=2)
        bad = self.root / 'bad.mp3'
        bad.write_bytes(b'not audio')
        with self.assertRaisesRegex(ValueError, '切片失败'):
            audio_pipeline.transcribe(bad, ffmpeg=ffmpeg, cache_dir=self.root / 'cache',
                cache_key='test', recognize=lambda data: '', check_active=lambda: None)


class DocumentTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def test_image_and_scanned_pdf_use_actual_offline_ocr(self):
        from PIL import Image, ImageDraw, ImageFont
        image = Image.new('RGB', (1200, 260), 'white')
        font_path = next((p for p in ('C:/Windows/Fonts/arial.ttf', '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf') if Path(p).exists()), None)
        font = ImageFont.truetype(font_path, 42) if font_path else ImageFont.load_default(size=42)
        ImageDraw.Draw(image).text((40, 60), 'Customer plan 2026', font=font, fill='black')
        path = self.root / 'scan.png'
        image.save(path)
        self.assertIn('Customer', document_ingest.extract(path)[0])
        scan = self.root / 'scan.pdf'
        image.save(scan, 'PDF', resolution=144)
        self.assertIn('Customer', document_ingest.extract(scan)[0])
        from pypdf import PdfReader, PdfWriter
        writer = PdfWriter()
        writer.add_blank_page(100, 100)
        writer.add_page(PdfReader(str(scan)).pages[0])
        mixed = self.root / 'mixed.pdf'
        writer.write(mixed)
        self.assertIn('Customer', document_ingest.extract(mixed)[0])

    def test_pptx_text_table_notes_and_empty_slide(self):
        from pptx import Presentation
        from pptx.util import Inches
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = '客户方案'
        table = slide.shapes.add_table(1, 2, Inches(1), Inches(2), Inches(5), Inches(1)).table
        table.cell(0, 0).text = '预算'
        table.cell(0, 1).text = '80000'
        slide.notes_slide.notes_text_frame.text = '审批人周五确认'
        path = self.root / 'plan.pptx'
        presentation.save(path)
        text = document_ingest.extract(path)[0]
        for expected in ('客户方案', '预算', '80000', '审批人周五确认'):
            self.assertIn(expected, text)
        blank = Presentation()
        blank.slides.add_slide(blank.slide_layouts[6])
        blank.save(path)
        self.assertEqual(document_ingest.extract(path)[0], '')

    def test_xls_and_binary_word_ppt_readers(self):
        import xlwt  # Test fixture writer only; not a runtime dependency.
        book = xlwt.Workbook()
        sheet = book.add_sheet('旧版报价')
        sheet.write(0, 0, '客户甲')
        sheet.write(0, 1, 80000)
        path = self.root / 'quote.xls'
        book.save(str(path))
        self.assertIn('客户甲', document_ingest.extract(path)[0])
        document = bytearray(1024)
        struct.pack_into('<H', document, 0, 0xA5EC)
        text = '客户需求\r报价确认'
        data = text.encode('utf-16le')
        document[600:600 + len(data)] = data
        plc = struct.pack('<II', 0, len(text)) + struct.pack('<HIH', 0, 600, 0)
        table = b'\x02' + struct.pack('<I', len(plc)) + plc
        struct.pack_into('<II', document, 0x1A2, 0, len(table))
        self.assertIn('报价确认', legacy_office.word_text(document, table))
        with self.assertRaises(ValueError):
            legacy_office.word_text(document, b'\x02\xff')
        struct.pack_into('<H', document, 10, 0x100)
        with self.assertRaisesRegex(ValueError, '加密'):
            legacy_office.word_text(document, table)
        atom = struct.pack('<HHI', 0, 4000, len(data)) + data
        container = struct.pack('<HHI', 15, 1006, len(atom)) + atom
        self.assertIn('报价确认', legacy_office.powerpoint_text(container))
        with self.assertRaisesRegex(ValueError, '损坏'):
            legacy_office.powerpoint_text(container[:-1])

    def test_real_legacy_files_exclude_ppt_masters_and_old_saves(self):
        fixtures = Path(__file__).parent / 'fixtures'
        self.assertIn('Word 97-SR2', document_ingest.extract(fixtures / 'simple.doc')[0])
        text = document_ingest.extract(fixtures / 'basic_test_ppt_file.ppt')[0]
        self.assertEqual(text.count('This is a test title'), 1)
        self.assertIn('This is page two', text)
        self.assertIn('These are the notes for page 1', text)
        self.assertNotIn('Click to edit Master', text)


if __name__ == '__main__':
    unittest.main()
