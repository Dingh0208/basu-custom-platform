"""Run with: python -m unittest discover -s services/ai -p test_sales.py -v"""

import io
import json
import os
import sqlite3
import tempfile
import threading
import unittest
from email.message import EmailMessage
from unittest.mock import patch
import urllib.error
import urllib.request
from urllib.parse import quote
import zipfile
from datetime import date, datetime, timedelta
from pathlib import Path

# Never load or mutate the user's application database.
_sandbox = tempfile.TemporaryDirectory(prefix="clienthub-sales-tests-")
os.environ["CLIENTHUB_DATA_DIR"] = _sandbox.name
import main
import sales
import document_ingest
from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter


class SalesTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        main.init_db()
        cls.server = main.ThreadingHTTPServer(("127.0.0.1", 0), main.Handler)
        cls.thread = threading.Thread(target=cls.server.serve_forever, daemon=True)
        cls.thread.start()
        cls.base = f"http://127.0.0.1:{cls.server.server_port}"

    @classmethod
    def tearDownClass(cls):
        cls.server.shutdown()
        cls.server.server_close()
        cls.thread.join()
        _sandbox.cleanup()

    def setUp(self):
        main.init_db()
        with main.db() as conn:
            for table in ("todos", "requirements", "projects", "customers", "imports", "chats", "tasks", "customer_match_audits"):
                conn.execute(f"DELETE FROM {table}")
            conn.execute("INSERT OR IGNORE INTO transcripts(id,title,file_name,file_size,duration,imported_at,customer,company,project,source,segments_json,summary_json) VALUES ('tr-001','测试通话','test.m4a','1 MB',60,'2026-10-20','测试客户','测试公司','测试项目','audio','[]','{}')")
            conn.execute("UPDATE transcripts SET customer='测试客户',project='测试项目',imported_at='2026-10-20 10:00',summary_json=?",
                         (json.dumps({"needs": "统一检索需求", "oneLine": "销售测试", "nextActions": "发送报价", "objections": "预算需要审批"}, ensure_ascii=False),))
        main.update_customer("测试客户", {"owner": "负责人甲", "company": "统一检索公司"})
        self.project = main.create_project("测试客户", {"name": "测试项目", "owner": "负责人甲"})
        self.deadline = (date.today() - timedelta(days=1)).isoformat()

    def request(self, path, method="GET", data=None):
        body = json.dumps(data).encode() if data is not None else None
        request = urllib.request.Request(self.base + path, data=body, method=method, headers={"Content-Type": "application/json"})
        try:
            with urllib.request.urlopen(request, timeout=10) as response:
                return response.status, json.loads(response.read())
        except urllib.error.HTTPError as exc:
            return exc.code, json.loads(exc.read())

    def requirement(self, title="统一检索需求", **extra):
        return main.create_requirement("测试客户", {"project": "测试项目", "title": title, "owner": "负责人甲", "dueDate": self.deadline, "sourceType": "audio", "sourceId": "tr-001", "sourceQuote": "原始需求证据", **extra})

    def todo(self, **extra):
        with main.db() as conn:
            return sales.save_todo(conn, {"customer": "测试客户", "project": "测试项目", "title": "发送报价", "owner": "负责人甲", "dueDate": self.deadline, "sourceId": "tr-001", "sourceType": "audio", **extra})

    def import_record(self, record_id, status="inbox", customer="待归档"):
        with main.db() as conn:
            conn.execute("INSERT INTO imports(id,file_name,type,size,customer,project,status,progress,imported_at) VALUES (?,?, 'PDF','1 KB',?,'—',?,100,'2026-10-20 10:00')", (record_id, record_id + '.pdf', customer, status))

    def test_migration_keeps_legacy_data_and_is_idempotent(self):
        conn = sqlite3.connect(":memory:")
        conn.row_factory = sqlite3.Row
        conn.executescript("CREATE TABLE imports(id TEXT); INSERT INTO imports VALUES ('legacy'); CREATE TABLE requirements(id TEXT);")
        sales.migrate(conn)
        sales.migrate(conn)
        self.assertEqual(conn.execute("SELECT id,owner FROM imports").fetchone()["id"], "legacy")
        self.assertEqual(conn.execute("SELECT owner FROM imports").fetchone()[0], "")
        conn.close()

    def test_attachment_migration_adds_transcript_column_once(self):
        conn = sqlite3.connect(':memory:')
        conn.row_factory = sqlite3.Row
        conn.executescript("CREATE TABLE chats(id TEXT); CREATE TABLE transcripts(id TEXT);")
        main.migrate_ingest(conn)
        main.migrate_ingest(conn)
        for table in ('chats', 'transcripts'):
            columns = {row['name'] for row in conn.execute(f'PRAGMA table_info({table})')}
            self.assertIn('attachments_json', columns)
        conn.close()

    def test_requirement_todo_trace_and_completion_persist(self):
        req = self.requirement()
        status, todo = self.request('/api/todos', 'POST', {"requirementId": req["id"], "title": req["title"]})
        self.assertEqual(status, 201)
        self.assertEqual((todo['sourceId'], todo['sourceQuote'], todo['owner'], todo['dueDate']), ('tr-001', '原始需求证据', '负责人甲', self.deadline))
        status, completed = self.request('/api/todos/' + todo['id'], 'PATCH', {"status": "done"})
        self.assertEqual(status, 200)
        self.assertIsNotNone(completed['completedAt'])
        self.assertEqual(self.request('/api/todos')[1][0]['status'], 'done')
        # Completing a follow-up is independent from implementing a requirement.
        self.assertEqual(main.list_requirements('测试客户')[0]['status'], '待确认')

    def test_duplicate_todo_does_not_create_second_record(self):
        self.assertEqual(self.todo()['id'], self.todo()['id'])
        self.assertEqual(len(self.request('/api/todos')[1]), 1)

    def test_validation_rejects_invalid_or_cross_customer_links(self):
        for extra in ({'dueDate': 'next week'}, {'owner': ''}, {'customer': '不存在'}, {'project': '其他客户的项目'}, {'sourceId': 'missing'}, {'status': 'invalid'}):
            payload = {'customer': '测试客户', 'project': '测试项目', 'title': '测试', 'owner': '甲', 'dueDate': self.deadline, **extra}
            self.assertEqual(self.request('/api/todos', 'POST', payload)[0], 400)

    def test_archive_is_atomic_and_updates_source_requirements(self):
        self.import_record('pending', 'done')
        self.assertEqual(main.dashboard_summary()['stats']['inbox'], 1)
        with main.db() as conn:
            conn.execute("UPDATE imports SET transcript_id='tr-001' WHERE id='pending'")
            conn.execute("UPDATE transcripts SET customer='待归档',project='—' WHERE id='tr-001'")
        main.create_requirement('待归档', {'title': '待归档需求', 'sourceId': 'tr-001'})
        payload = {'ids': ['pending', 'missing'], 'customer': '测试客户', 'project': '测试项目', 'owner': '负责人乙'}
        self.assertEqual(self.request('/api/imports/archive', 'POST', payload)[0], 400)
        with main.db() as conn:
            self.assertEqual(conn.execute("SELECT customer FROM imports WHERE id='pending'").fetchone()[0], '待归档')
        payload['ids'] = ['pending']
        self.assertEqual(self.request('/api/imports/archive', 'POST', payload)[0], 200)
        req = main.list_requirements('测试客户')[0]
        self.assertEqual((req['project'], req['owner']), ('测试项目', '负责人乙'))
        self.assertEqual(self.request('/api/transcripts/tr-001')[1]['customer'], '测试客户')
        self.assertEqual(main.dashboard_summary()['stats']['inbox'], 0)

    def test_archive_rejects_processing_and_wrong_project(self):
        self.import_record('pending', 'processing')
        payload = {'ids': ['pending'], 'customer': '测试客户', 'project': '测试项目', 'owner': '甲'}
        self.assertEqual(self.request('/api/imports/archive', 'POST', payload)[0], 400)
        payload['project'] = '不存在'
        self.assertEqual(self.request('/api/imports/archive', 'POST', payload)[0], 400)

    def test_archive_shared_source_requires_all_affected_imports(self):
        for record_id in ('pending-a', 'pending-b'):
            self.import_record(record_id)
        with main.db() as conn:
            conn.execute("UPDATE imports SET transcript_id='tr-001'")
        payload = {'ids': ['pending-a'], 'customer': '测试客户', 'project': '测试项目', 'owner': '甲'}
        self.assertEqual(self.request('/api/imports/archive', 'POST', payload)[0], 400)
        payload['ids'].append('pending-b')
        self.assertEqual(self.request('/api/imports/archive', 'POST', payload)[0], 200)

    def test_rename_customer_and_project_preserves_todo(self):
        req = self.requirement()
        todo = self.todo(requirementId=req['id'])
        main.update_customer('测试客户', {'name': '新客户名'})
        main.update_project(self.project['id'], {'name': '新项目名'})
        saved = self.request('/api/todos')[1][0]
        self.assertEqual((saved['customer'], saved['project'], saved['id']), ('新客户名', '新项目名', todo['id']))
        main.update_requirement(req['id'], {'owner': '新负责人', 'dueDate': '2030-01-01'})
        saved = self.request('/api/todos')[1][0]
        self.assertEqual((saved['owner'], saved['dueDate']), ('新负责人', '2030-01-01'))

    def test_merge_requirement_rebinds_and_delete_preserves_todo(self):
        first, second = self.requirement('需求一', sourceType='manual', sourceId=''), self.requirement('需求二')
        self.todo(requirementId=second['id'])
        main.merge_requirements(first['id'], [second['id']])
        todo = self.request('/api/todos')[1][0]
        self.assertEqual(todo['requirementId'], first['id'])
        completed = self.request('/api/todos/' + todo['id'], 'PATCH', {'status': 'done'})[1]
        self.assertEqual(completed['sourceId'], 'tr-001')
        main.delete_requirement(first['id'])
        todo = self.request('/api/todos')[1][0]
        self.assertEqual(todo['requirementId'], '')
        self.assertEqual(todo['sourceId'], 'tr-001')

    def test_requirement_description_edit_preserves_todo_assignment(self):
        req = self.requirement()
        self.todo(requirementId=req['id'], owner='负责执行的人', dueDate='2030-01-01')
        main.update_requirement(req['id'], {'description': '补充说明'})
        saved = self.request('/api/todos')[1][0]
        self.assertEqual((saved['owner'], saved['dueDate']), ('负责执行的人', '2030-01-01'))

    def test_project_and_customer_delete_do_not_leave_orphans(self):
        self.todo()
        main.delete_project(self.project['id'])
        self.assertEqual(self.request('/api/todos')[1][0]['project'], '—')
        main.delete_customer('测试客户')
        self.assertEqual(self.request('/api/todos')[1], [])

    def test_delete_customer_removes_projects_and_does_not_reseed_on_restart(self):
        self.requirement()
        self.todo()
        self.import_record('owned', 'done', '测试客户')
        with main.db() as conn:
            conn.execute("INSERT INTO customer_match_audits(id,import_id,method,created_at) VALUES ('match-owned','owned','manual','2026-10-30')")
        result = main.delete_customer('测试客户')
        self.assertEqual(result['projects'], 1)
        self.assertEqual(main.list_projects('测试客户'), [])
        main.init_db()
        main.backfill_requirements()
        self.assertEqual(main.list_all_projects(), [])
        self.assertEqual(main.list_customers(), [])
        with main.db() as conn:
            for table in ('imports', 'transcripts', 'requirements', 'todos', 'customer_match_audits'):
                self.assertEqual(conn.execute(f'SELECT COUNT(*) FROM {table}').fetchone()[0], 0)

    def test_profile_is_validated_and_persistent(self):
        with main.db() as conn:
            conn.execute("DELETE FROM settings WHERE key LIKE 'profile.%'")
        self.assertEqual(self.request('/api/profile')[1]['name'], '王总')
        self.assertEqual(self.request('/api/profile', 'POST', {'name': ' '})[0], 400)
        self.assertEqual(self.request('/api/profile', 'POST', {'name': '李明', 'phone': []})[0], 400)
        payload = {'name': '李明', 'role': '销售经理', 'company': '示例公司', 'phone': '13800000000', 'email': 'li@example.com'}
        self.assertEqual(self.request('/api/profile', 'POST', payload), (200, payload))
        main.init_db()
        self.assertEqual(self.request('/api/profile')[1], payload)

    def test_sales_endpoints_are_available_before_any_todos_exist(self):
        from urllib.parse import quote
        self.assertEqual(self.request('/api/todos'), (200, []))
        status, followup = self.request('/api/customers/' + quote('测试客户') + '/followup')
        self.assertEqual(status, 200)
        self.assertEqual(followup['todos'], [])
        metrics_status, metrics = self.request('/api/metrics?days=7')
        self.assertEqual(metrics_status, 200)
        self.assertIn('importsSuccessRate', metrics)
        dashboard_status, dashboard = self.request('/api/dashboard')
        self.assertEqual(dashboard_status, 200)
        self.assertIn('metrics', dashboard)
        self.assertGreaterEqual(self.request('/api/health')[1]['apiRevision'], 4)

    def test_deleting_one_shared_import_keeps_other_source_todos(self):
        for record_id in ('audio-a', 'audio-b'):
            self.import_record(record_id, 'done', '测试客户')
        with main.db() as conn:
            conn.execute("UPDATE imports SET transcript_id='tr-001'")
        self.todo(requirementId=self.requirement()['id'])
        main.delete_import('audio-a')
        self.assertEqual(len(self.request('/api/todos')[1]), 1)
        main.delete_import('audio-b')
        self.assertEqual(self.request('/api/todos')[1], [])

    def test_customer_merge_preserves_project_and_todo_links(self):
        self.todo(requirementId=self.requirement()['id'])
        main.update_customer('目标客户', {})
        main.create_project('目标客户', {'name': '测试项目'})
        main.merge_customers('测试客户', '目标客户')
        todo = self.request('/api/todos')[1][0]
        self.assertEqual(todo['customer'], '目标客户')
        self.assertEqual(main.list_requirements('目标客户')[0]['id'], todo['requirementId'])
        self.assertEqual(self.request('/api/todos/' + todo['id'], 'PATCH', {'status': 'done'})[0], 200)
        with main.db() as conn:
            self.assertEqual(conn.execute("SELECT COUNT(*) FROM projects WHERE customer='目标客户' AND name='测试项目'").fetchone()[0], 1)

    def test_followup_tracks_due_risks_and_hides_confirmed_suggestions(self):
        with main.db() as conn:
            data = sales.followup(conn, '测试客户')
        self.assertEqual(data['suggestions'][0]['title'], '发送报价')
        self.assertEqual(data['risks'][0]['text'], '预算需要审批')
        self.todo()
        with main.db() as conn:
            data = sales.followup(conn, '测试客户')
        self.assertEqual(data['overdue'], 1)
        self.assertEqual(data['suggestions'], [])
        self.assertEqual(data['lastContactAt'], '2026-10-20')

    def test_followup_recognizes_import_and_transcript_as_same_source(self):
        self.import_record('audio-import', 'done', '测试客户')
        with main.db() as conn:
            conn.execute("UPDATE imports SET transcript_id='tr-001' WHERE id='audio-import'")
        req = self.requirement(title='发送报价', sourceId='audio-import')
        self.todo(requirementId=req['id'])
        with main.db() as conn:
            self.assertEqual(sales.followup(conn, '测试客户')['suggestions'], [])

    def test_chat_role_change_moves_linked_sales_records(self):
        with main.db() as conn:
            conn.execute("INSERT INTO chats(id,title,file_name,customer,project,imported_at,participants_json,messages_json,analysis_json,status,message_count) VALUES ('chat-test','聊天','chat.txt','测试客户','测试项目','2026-10-30',?,'[]','{}','done',1)",
                         (json.dumps([{'name': '测试客户', 'role': 'customer'}, {'name': '新客户', 'role': 'me'}]),))
        req = self.requirement(sourceType='chat', sourceId='chat-test')
        self.todo(requirementId=req['id'])
        main.update_chat_participants('chat-test', [{'name': '测试客户', 'role': 'me'}, {'name': '新客户', 'role': 'customer'}])
        self.assertEqual(main.list_requirements('新客户')[0]['id'], req['id'])
        saved = self.request('/api/todos')[1][0]
        self.assertEqual(saved['customer'], '新客户')
        self.assertEqual(self.request('/api/todos/' + saved['id'], 'PATCH', {'status': 'done'})[0], 200)

    def test_followup_prefers_chat_date_over_later_import_date(self):
        with main.db() as conn:
            conn.execute("INSERT INTO chats(id,title,file_name,customer,project,imported_at,ended_at,participants_json,messages_json,analysis_json,status,message_count) VALUES ('chat-test','聊天','chat.txt','测试客户','测试项目','2026-10-30','2026-10-21 14:00','[]','[]','{}','done',1)")
            self.assertEqual(sales.followup(conn, '测试客户')['lastContactAt'], '2026-10-21')

    def test_search_five_entities_and_no_stale_results(self):
        self.requirement()
        main.update_project(self.project['id'], {'description': '统一检索方案'})
        with main.db() as conn:
            conn.execute("INSERT INTO chats(id,title,file_name,customer,project,imported_at,participants_json,messages_json,analysis_json,status,message_count) VALUES ('chat-test','聊天','chat.txt','测试客户','测试项目','2026-10-20','[]',?,'{}','done',1)", (json.dumps([{'content': '统一检索聊天'}], ensure_ascii=False),))
            result = sales.search(conn, '统一检索')
            literal = sales.search(conn, '%')
        self.assertEqual({r['kind'] for r in result['results']}, {'customer', 'project', 'requirement', 'chat', 'transcript'})
        self.assertEqual(literal['results'], [])
        main.delete_requirement(main.list_requirements('测试客户')[0]['id'])
        with main.db() as conn:
            self.assertEqual(sales.search(conn, '统一检索')['counts']['requirement'], 0)

    def test_summary_patch_is_persistent_and_validated(self):
        self.assertEqual(self.request('/api/transcripts/tr-001/summary', 'PATCH', {'summary': {'needs': []}})[0], 400)
        self.assertEqual(self.request('/api/transcripts/tr-001/summary', 'PATCH', {'summary': {'oneLine': '修改后的摘要'}})[0], 200)
        summary = self.request('/api/transcripts/tr-001')[1]['summary']
        self.assertEqual(summary['oneLine'], '修改后的摘要')
        self.assertEqual(summary['needs'], '统一检索需求')

    def test_search_caps_results_but_counts_all_matches(self):
        for number in range(25):
            self.requirement(title=f'报价跟进 {number}')
        with main.db() as conn:
            result = sales.search(conn, '报价跟进')
        self.assertEqual(result['counts']['requirement'], 25)
        self.assertEqual(len(result['results']), 20)

    def test_product_metrics_report_rates_latency_and_unknown_samples(self):
        anchor = datetime(2026, 10, 31, 12, 0, 0)
        self.import_record('metric-done', 'done', '测试客户')
        self.import_record('metric-failed', 'failed', '测试客户')
        with main.db() as conn:
            conn.execute("UPDATE imports SET imported_at='2026-10-30 09:00' WHERE id IN ('metric-done','metric-failed')")
            conn.execute("UPDATE transcripts SET imported_at='2026-10-29 09:00' WHERE id='tr-001'")
            conn.executemany(
                "INSERT INTO tasks(id,kind,status,created_at,updated_at,finished_at) VALUES (?,?,?,?,?,?)",
                [
                    ('metric-task-ok', 'import_real', 'succeeded', '2026-10-30 10:00:00', '2026-10-30 10:01:00', '2026-10-30 10:01:00'),
                    ('metric-task-fail', 'import_document', 'failed', '2026-10-30 10:00:00', '2026-10-30 10:10:00', '2026-10-30 10:10:00'),
                ],
            )
            conn.execute("INSERT INTO customer_match_audits(id,import_id,suggested_customer,confirmed_customer,method,is_correct,created_at,confirmed_at) VALUES ('match-ok','metric-done','测试客户','测试客户','auto',1,'2026-10-30','2026-10-30')")
            conn.execute("INSERT INTO customer_match_audits(id,import_id,suggested_customer,confirmed_customer,method,is_correct,created_at,confirmed_at) VALUES ('match-bad','metric-failed','错误客户','测试客户','auto',0,'2026-10-30','2026-10-30')")
        self.requirement(title='已确认需求', status='已确认')
        todo = self.todo()
        with main.db() as conn:
            conn.execute("UPDATE todos SET created_at='2026-10-30 09:00' WHERE id=?", (todo['id'],))
            report = sales.metrics(conn, days=7, now=anchor)
        self.assertEqual(report['importsSuccessRate']['value'], 0.5)
        self.assertEqual(report['parseCompletionRate']['value'], 0.5)
        self.assertEqual(report['parseP95Seconds']['value'], 600.0)
        self.assertEqual(report['customerMatchAccuracy']['value'], 0.5)
        self.assertEqual(report['requirementEffectiveness']['value'], 1.0)
        self.assertEqual(report['nextActionLatencyMinutes']['value'], 1440.0)
        self.assertEqual(report['weeklyActiveCustomers']['value'], 1)
        self.assertEqual(report['weeklyActiveProjects']['value'], 1)

    def test_metrics_leave_unmeasured_accuracy_as_null(self):
        with main.db() as conn:
            report = sales.metrics(conn, days=7, now=datetime(2026, 10, 31, 12, 0, 0))
        self.assertIsNone(report['customerMatchAccuracy']['value'])
        self.assertFalse(report['customerMatchAccuracy']['available'])

    def test_backup_round_trip_includes_sales_todos(self):
        todo = self.todo()
        archive, _ = main.create_backup_archive()
        with zipfile.ZipFile(io.BytesIO(archive)) as zipped:
            self.assertIn('clienthub.db', zipped.namelist())
        with main.db() as conn:
            conn.execute('DELETE FROM todos')
        main.restore_backup_archive(archive)
        self.assertEqual(self.request('/api/todos')[1][0]['id'], todo['id'])

    def test_document_parsers_extract_office_text_and_eml_attachments(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            docx_path = root / 'brief.docx'
            document = Document()
            document.add_paragraph('Document body')
            document.save(docx_path)
            self.assertIn('Document body', document_ingest.extract(docx_path)[0])

            xlsx_path = root / 'brief.xlsx'
            workbook = Workbook()
            workbook.active.title = 'Pipeline'
            workbook.active.append(['客户', '下次跟进'])
            workbook.active.append(['新客户', '周五'])
            workbook.save(xlsx_path)
            self.assertIn('Pipeline', document_ingest.extract(xlsx_path)[0])
            self.assertIn('新客户', document_ingest.extract(xlsx_path)[0])

            mail = EmailMessage()
            mail['From'] = 'sender@example.test'
            mail['To'] = 'sales@example.test'
            mail['Subject'] = 'Follow up'
            mail.set_content('Please review the attached plan.')
            mail.add_attachment(docx_path.read_bytes(), maintype='application', subtype='vnd.openxmlformats-officedocument.wordprocessingml.document', filename='plan.docx')
            eml_path = root / 'follow-up.eml'
            eml_path.write_bytes(mail.as_bytes())
            text, attachments = document_ingest.extract(eml_path)
            self.assertIn('Please review the attached plan.', text)
            self.assertEqual(attachments[0][0], 'plan.docx')
            self.assertEqual(attachments[0][1], docx_path.read_bytes())

            from pypdf import PdfWriter
            writer = PdfWriter()
            writer.add_blank_page(width=100, height=100)
            pdf_path = root / 'scan.pdf'
            with pdf_path.open('wb') as output:
                writer.write(output)
            with self.assertRaisesRegex(ValueError, 'OCR'):
                document_ingest.extract(pdf_path)

    def test_office_original_and_mail_attachments_are_persisted_and_previewable(self):
        mail = EmailMessage()
        mail['Subject'] = 'Customer update'
        mail.set_content('The rollout starts next month.')
        mail.add_attachment(b'Quarterly plan', maintype='text', subtype='plain', filename='plan.txt')
        import_id = 'email-preview'
        source_path = main.UPLOAD_DIR / f'{import_id}.eml'
        source_path.write_bytes(mail.as_bytes())
        with main.db() as conn:
            conn.execute("INSERT INTO imports(id,file_name,type,size,customer,project,status,progress,imported_at) VALUES (?, 'follow-up.eml', '邮件', '1 KB', '测试客户', '测试项目', 'processing', 0, '2026-10-20')", (import_id,))
        with patch('main.get_llm_config', return_value={'apiKey': ''}):
            main.process_document(import_id, source_path)
        transcript = self.request('/api/transcripts/tr-email-preview')[1]
        self.assertEqual(len(transcript['attachments']), 2)
        source, attachment = transcript['attachments']
        self.assertTrue((main.UPLOAD_DIR / source['storedName']).is_file())
        preview_status, preview = self.request(f"/api/files/{source['storedName']}/preview")
        self.assertEqual(preview_status, 200)
        self.assertIn('The rollout starts next month.', preview['text'])
        preview_status, preview = self.request(f"/api/files/{attachment['storedName']}/preview")
        self.assertEqual(preview_status, 200)
        self.assertEqual(preview['text'], 'Quarterly plan')
        self.assertEqual(transcript['summary']['confidence'], 0)
        archive, _ = main.create_backup_archive()
        main.restore_backup_archive(archive)
        self.assertEqual((main.UPLOAD_DIR / source['storedName']).read_bytes(), mail.as_bytes())
        self.assertEqual(self.request('/api/transcripts/tr-email-preview')[1]['attachments'][1]['storedName'], attachment['storedName'])

    def test_voice_transcription_uses_persistent_task_and_updates_chat(self):
        source_id = 'chat-voice-test'
        stored_name = f'{source_id}.mp3'
        (main.UPLOAD_DIR / stored_name).write_bytes(b'test audio')
        attachments = [{'id': 'voice-1', 'kind': 'voice', 'name': 'voice.mp3', 'storedName': stored_name}]
        with main.db() as conn:
            conn.execute("INSERT INTO chats(id,title,file_name,customer,project,imported_at,participants_json,messages_json,analysis_json,status,message_count,started_at,ended_at,attachments_json) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)", (source_id, 'Test', 'chat.txt', '测试客户', '测试项目', '2026-10-20', '[]', '[]', '{}', 'done', 0, '', '', json.dumps(attachments)))
        with patch('main.transcribe_audio_file', return_value='已确认下周进行评审'), patch('main.start_task', side_effect=lambda task_id, runner, *args: main._run_task(task_id, runner, *args)):
            status, response = self.request(f'/api/chats/{source_id}/attachments/transcribe', 'POST', {'attachmentId': 'voice-1'})
        self.assertEqual(status, 202)
        self.assertEqual(main.get_task(response['taskId'])['status'], 'succeeded')
        record = self.request(f'/api/chats/{source_id}')[1]
        self.assertEqual(record['attachments'][0]['transcript'], '已确认下周进行评审')
        self.assertEqual(record['attachments'][0]['transcriptStatus'], 'done')
        with main.db() as conn:
            self.assertEqual(sales.search(conn, '已确认下周进行评审')['counts']['chat'], 1)

        transcript_audio = 'transcript-voice-test.mp3'
        (main.UPLOAD_DIR / transcript_audio).write_bytes(b'test audio')
        with main.db() as conn:
            conn.execute("UPDATE transcripts SET attachments_json=? WHERE id='tr-001'", (json.dumps([{'id': 'voice-2', 'kind': 'voice', 'name': transcript_audio, 'storedName': transcript_audio}], ensure_ascii=False),))
        with patch('main.transcribe_audio_file', return_value='报价已确认'), patch('main.start_task', side_effect=lambda task_id, runner, *args: main._run_task(task_id, runner, *args)):
            status, response = self.request('/api/transcripts/tr-001/attachments/transcribe', 'POST', {'attachmentId': 'voice-2'})
        self.assertEqual(status, 202)
        self.assertEqual(self.request('/api/transcripts/tr-001')[1]['attachments'][0]['transcript'], '报价已确认')

    def test_segment_corrections_are_atomic_and_invalidate_summary_evidence(self):
        with main.db() as conn:
            conn.execute("UPDATE transcripts SET segments_json=?,summary_json=? WHERE id='tr-001'", (json.dumps([{'id': 's1', 'start': 0, 'end': 1, 'speaker': 'unknown', 'text': '原始内容'}], ensure_ascii=False), json.dumps({'needs': '原始内容', 'evidence': {'needs': '原始内容'}, 'confidence': 1, 'evidenceCount': 1}, ensure_ascii=False)))
        invalid = [{'id': 's1', 'speaker': 'customer', 'text': '已修改'}, {'id': 'missing', 'speaker': 'me', 'text': ''}]
        self.assertEqual(self.request('/api/transcripts/tr-001/segments', 'PATCH', {'segments': invalid})[0], 400)
        unchanged = self.request('/api/transcripts/tr-001')[1]
        self.assertEqual(unchanged['segments'][0]['text'], '原始内容')
        self.assertEqual(unchanged['corrections'], [])
        valid = [{'id': 's1', 'speaker': 'customer', 'text': '已修改'}]
        status, updated = self.request('/api/transcripts/tr-001/segments', 'PATCH', {'segments': valid})
        self.assertEqual(status, 200)
        self.assertEqual(updated['summary']['confidence'], 0)
        self.assertEqual(updated['summary'].get('evidence'), None)
        self.assertEqual(updated['corrections'][0]['after'], '已修改')

    def test_unverified_ai_evidence_does_not_create_requirements(self):
        response = {'choices': [{'message': {'content': json.dumps({'needs': '新增模块', 'evidence': {'needs': '模型编造的原文'}})}}]}
        with patch('main.maas_request', return_value=response):
            summary = main.summarize_transcript('实际通话没有需求')
        self.assertEqual(summary['confidence'], 0)
        self.assertEqual(summary['evidence'], {})
        main.insert_requirements_from_audio('imp-fake', '测试客户', '测试项目', summary)
        self.assertEqual(self.request(f'/api/customers/{quote("测试客户")}/requirements')[1], [])

    def test_regenerate_summary_uses_saved_text_and_persists_verified_result(self):
        with main.db() as conn:
            conn.execute("UPDATE transcripts SET segments_json=? WHERE id='tr-001'", (json.dumps([{'id': 's1', 'text': '客户明确需要报价'}], ensure_ascii=False),))
        generated = {'needs': '报价', 'evidence': {'needs': '客户明确需要报价'}, 'confidence': 1, 'evidenceCount': 1}
        with patch('main.get_llm_config', return_value={'apiKey': 'test-key'}), patch('main.summarize_transcript', return_value=generated) as summarize:
            status, record = self.request('/api/transcripts/tr-001/regenerate-summary', 'POST', {})
        self.assertEqual(status, 200)
        summarize.assert_called_once_with('客户明确需要报价')
        self.assertEqual(record['summary']['evidence']['needs'], '客户明确需要报价')

    def test_group_role_edit_keeps_the_selected_customer(self):
        source_id = 'chat-group-test'
        participants = [{'name': '销售', 'role': 'me'}, {'name': '客户甲', 'role': 'customer'}, {'name': '客户乙', 'role': 'customer'}]
        with main.db() as conn:
            conn.execute("INSERT INTO chats(id,title,file_name,customer,project,imported_at,participants_json,messages_json,analysis_json,status,message_count,started_at,ended_at,attachments_json) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)", (source_id, 'Group', 'chat.txt', '客户乙', '测试项目', '2026-10-20', json.dumps(participants), '[]', '{}', 'done', 0, '', '', '[]'))
        updated = main.update_chat_participants(source_id, participants)
        self.assertEqual(updated['customer'], '客户乙')
        self.assertEqual(len(updated['participants']), 3)


if __name__ == '__main__':
    unittest.main()
