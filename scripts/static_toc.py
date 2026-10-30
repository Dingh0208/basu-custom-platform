from docx import Document
from docx.enum.text import WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

DOCX = r"D:\code\Basu-Custom-Platform\docs\PRD-客户信息聚合工作台-v1.0.docx"


def parse_entry(text):
    text = text.replace("\u00a0", " ")
    if "\t" in text:
        title, page = text.rsplit("\t", 1)
        return title.strip(), page.strip()
    # Fallback: trailing digits are the page number.
    i = len(text) - 1
    while i >= 0 and text[i].isdigit():
        i -= 1
    return text[:i + 1].strip(), text[i + 1:].strip()


def main():
    doc = Document(DOCX)
    toc_paras = [p for p in doc.paragraphs if p.style.name.lower() == "toc 2"]
    entries = [parse_entry(p.text) for p in toc_paras if p.text.strip()]
    print("entries:", len(entries))

    # Find the first content heading after the TOC.
    first_content = None
    seen_toc = False
    for p in doc.paragraphs:
        if p.style.name.lower() == "toc 2":
            seen_toc = True
            continue
        if seen_toc and p.style.name == "Heading 2":
            first_content = p
            break
    if first_content is None:
        raise RuntimeError("Could not find first content heading")

    # The Word-updated field leaves a page break before the TOC entries.
    # Remove it; the TOC should start on the page after the cover.
    paras = doc.paragraphs
    heading_idx = next(i for i, p in enumerate(paras) if p.style.name == "TOC Heading" and p.text.strip() == "目录")
    first_toc_idx = next(i for i, p in enumerate(paras) if p.style.name.lower() == "toc 2")
    for i in range(first_toc_idx - 1, heading_idx, -1):
        para = paras[i]
        brs = para._p.findall(".//" + qn("w:br"))
        if any(b.get(qn("w:type")) == "page" for b in brs):
            doc.element.body.remove(para._element)

    for p in toc_paras:
        p._element.getparent().remove(p._element)

    for title, page in entries:
        p = first_content.insert_paragraph_before()
        p.style = doc.styles["toc 2"]
        p.paragraph_format.tab_stops.add_tab_stop(
            Cm(16.3), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
        )
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.1
        r1 = p.add_run(title)
        r1.font.size = Pt(10.5)
        r1.font.name = "微软雅黑"
        r1._element.get_or_add_rPr().rFonts.set(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia",
            "微软雅黑",
        )
        r2 = p.add_run("\t")
        r2.font.size = Pt(10.5)
        r3 = p.add_run(page)
        r3.font.size = Pt(10.5)
        r3.font.name = "微软雅黑"

    # TOC page ends here; content starts on a new page.
    brk = first_content.insert_paragraph_before()
    brk.add_run().add_break(WD_BREAK.PAGE)

    doc.save(DOCX)
    print("static TOC written:", DOCX)


if __name__ == "__main__":
    main()
