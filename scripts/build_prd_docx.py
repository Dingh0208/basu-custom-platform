import os
import subprocess
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCS = os.path.join(BASE, "docs")
MD = os.path.join(DOCS, "PRD-客户信息聚合工作台-v1.0.md")
REF = os.path.join(DOCS, "reference.docx")
OUT = os.path.join(DOCS, "PRD-客户信息聚合工作台-v1.0.docx")
PANDOC = r"D:\develop\Anaconda3\Library\bin\pandoc.exe"


def set_eastasia(style, font_name="微软雅黑"):
    style.font.name = font_name
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), font_name)
    rfonts.set(qn("w:ascii"), font_name)
    rfonts.set(qn("w:hAnsi"), font_name)


def build_reference():
    # Get pandoc's default reference docx (binary) to inherit its table/list styles.
    with open(REF, "wb") as f:
        subprocess.run([PANDOC, "--print-default-data-file", "reference.docx"], stdout=f, check=True)

    doc = Document(REF)

    # Page setup: A4, moderate margins.
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.2)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)
        section.header_distance = Cm(1.0)
        section.footer_distance = Cm(1.0)

    styles = doc.styles

    normal = styles["Normal"]
    set_eastasia(normal)
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.space_before = Pt(0)

    for name in ["Body Text", "First Paragraph"]:
        try:
            st = styles[name]
            set_eastasia(st)
            st.font.size = Pt(10.5)
            st.paragraph_format.line_spacing = 1.35
            st.paragraph_format.space_after = Pt(5)
        except KeyError:
            pass

    try:
        compact = styles["Compact"]
        set_eastasia(compact)
        compact.font.size = Pt(9.5)
    except KeyError:
        pass

    heading_specs = [
        ("Heading 1", 18, "1F4E79", 18, 10),
        ("Heading 2", 14, "2E75B6", 14, 7),
        ("Heading 3", 12, "333333", 10, 4),
    ]
    for name, size, color, before, after in heading_specs:
        try:
            st = styles[name]
            set_eastasia(st)
            st.font.size = Pt(size)
            st.font.bold = True
            st.font.color.rgb = RGBColor.from_string(color)
            st.paragraph_format.space_before = Pt(before)
            st.paragraph_format.space_after = Pt(after)
            st.paragraph_format.keep_with_next = True
        except KeyError:
            pass

    for name, size in [("Title", 30), ("Subtitle", 15)]:
        try:
            st = styles[name]
            set_eastasia(st)
            st.font.size = Pt(size)
            st.font.bold = name == "Title"
        except KeyError:
            pass

    # Header / footer. Cover page has no header/footer.
    section = doc.sections[0]
    section.different_first_page_header_footer = True
    hp = section.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hr = hp.add_run("客户信息聚合工作台（ClientHub）产品需求文档 · V1.0")
    hr.font.size = Pt(8.5)
    hr.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
    hr.font.name = "微软雅黑"

    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = fp.add_run("内部资料 · 第 ")
    r1.font.size = Pt(8.5)
    r1.font.name = "微软雅黑"
    add_field(fp, "PAGE", 8.5)
    r2 = fp.add_run(" 页 / 共 ")
    r2.font.size = Pt(8.5)
    r2.font.name = "微软雅黑"
    add_field(fp, "NUMPAGES", 8.5)
    r3 = fp.add_run(" 页")
    r3.font.size = Pt(8.5)
    r3.font.name = "微软雅黑"

    doc.save(REF)
    print("reference saved:", REF)


def add_field(paragraph, instr, size=9):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = instr
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    txt = OxmlElement("w:t")
    txt.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr_text)
    run._r.append(fld_sep)
    run._r.append(txt)
    run._r.append(fld_end)
    run.font.size = Pt(size)
    run.font.name = "微软雅黑"


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def format_tables(doc):
    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        tbl = table._tbl
        tblPr = tbl.tblPr
        # width 100%
        for old in tblPr.findall(qn("w:tblW")):
            tblPr.remove(old)
        tblW = OxmlElement("w:tblW")
        tblW.set(qn("w:w"), "5000")
        tblW.set(qn("w:type"), "pct")
        tblPr.append(tblW)
        # borders
        for old in tblPr.findall(qn("w:tblBorders")):
            tblPr.remove(old)
        borders = OxmlElement("w:tblBorders")
        for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), "BFBFBF")
            borders.append(el)
        tblPr.append(borders)

        for r_idx, row in enumerate(table.rows):
            trPr = row._tr.get_or_add_trPr()
            cantSplit = OxmlElement("w:cantSplit")
            trPr.append(cantSplit)
            if r_idx == 0:
                tblHeader = OxmlElement("w:tblHeader")
                tblHeader.set(qn("w:val"), "true")
                trPr.append(tblHeader)
            for cell in row.cells:
                if r_idx == 0:
                    shade(cell, "1F4E79")
                elif r_idx % 2 == 0:
                    shade(cell, "F2F6FA")
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(1)
                    p.paragraph_format.space_after = Pt(1)
                    p.paragraph_format.line_spacing = 1.0
                    for run in p.runs:
                        run.font.size = Pt(9)
                        run.font.name = "微软雅黑"
                        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "微软雅黑")
                        if r_idx == 0:
                            run.font.bold = True
                            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)


def remove_intro(doc):
    """Drop the Markdown H1 and metadata block; the cover carries them."""
    paragraphs = doc.paragraphs
    for i, p in enumerate(paragraphs[:6]):
        if p.style.name == "Heading 1" and "产品需求文档" in p.text:
            p._element.getparent().remove(p._element)
            # Remove the following metadata paragraph (blockquote collapsed into one).
            for q in doc.paragraphs[i:i+2]:
                if "文档版本" in q.text and "文档日期" in q.text:
                    q._element.getparent().remove(q._element)
                    break
            break


def add_toc(doc):
    first = doc.paragraphs[0]

    heading = first.insert_paragraph_before("目录")
    heading.style = doc.styles["TOC Heading"] if "TOC Heading" in [s.name for s in doc.styles] else doc.styles["Heading 1"]
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT

    p = first.insert_paragraph_before()
    run = p.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "2-2" \\h \\z \\u'
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    txt = OxmlElement("w:t")
    txt.text = "右键选择“更新域”生成目录"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(txt)
    run._r.append(fld_end)
    run.font.size = Pt(10.5)
    run.font.name = "微软雅黑"

    brk = first.insert_paragraph_before()
    brk.add_run().add_break(WD_BREAK.PAGE)


def add_cover(doc):
    first = doc.paragraphs[0]

    def add_before(text="", size=12, bold=False, color=None, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=0):
        p = first.insert_paragraph_before()
        p.alignment = align
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        if text:
            r = p.add_run(text)
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.name = "微软雅黑"
            r._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "微软雅黑")
            if color:
                r.font.color.rgb = RGBColor.from_string(color)
        return p

    add_before("内部资料 · 请勿外传", 9, False, "808080", WD_ALIGN_PARAGRAPH.RIGHT, 0, 0)
    for _ in range(4):
        add_before("", 10)
    add_before("客户信息聚合工作台", 34, True, "1F4E79", space_before=0, space_after=8)
    add_before("ClientHub", 20, False, "2E75B6", space_after=24)
    add_before("产品需求文档（PRD）", 22, True, "333333", space_after=40)
    add_before("版本：V1.0", 12, False, "595959", space_after=6)
    add_before("日期：2026 年 10 月 28 日", 12, False, "595959", space_after=6)
    add_before("适用形态：Windows 单机桌面应用 · 单用户", 12, False, "595959", space_after=6)
    add_before("编制：产品团队", 12, False, "595959", space_after=6)
    for _ in range(4):
        add_before("", 10)
    add_before("本文件用于产品需求评审与开发实现，未经允许请勿外传。", 9, False, "808080")
    p = add_before("", 10)
    p.add_run().add_break(WD_BREAK.PAGE)


def set_all_fonts(doc):
    for p in doc.paragraphs:
        for r in p.runs:
            if not r.font.name:
                r.font.name = "微软雅黑"
            r._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "微软雅黑")
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.name = "微软雅黑"
                        r._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "微软雅黑")


def fit_images(doc, max_width_cm=16.0):
    for shape in doc.inline_shapes:
        old_w = shape.width
        old_h = shape.height
        if old_w and old_h and old_w > Cm(max_width_cm):
            ratio = Cm(max_width_cm) / old_w
            shape.width = Cm(max_width_cm)
            shape.height = int(old_h * ratio)


def update_fields(doc):
    settings = doc.settings.element
    for old in settings.findall(qn("w:updateFields")):
        settings.remove(old)
    el = OxmlElement("w:updateFields")
    el.set(qn("w:val"), "true")
    settings.append(el)
    # LibreOffice shades field results by default; ask readers not to.
    for old in settings.findall(qn("w:doNotShadeFormData")):
        settings.remove(old)
    settings.append(OxmlElement("w:doNotShadeFormData"))


def main():
    build_reference()

    subprocess.run([
        PANDOC,
        MD,
        "-o", OUT,
        "--reference-doc", REF,
        "--from", "gfm",
        "--resource-path", DOCS,
        "--wrap=none",
    ], check=True)

    doc = Document(OUT)
    remove_intro(doc)
    add_toc(doc)
    add_cover(doc)
    format_tables(doc)
    fit_images(doc)
    set_all_fonts(doc)
    update_fields(doc)

    props = doc.core_properties
    props.title = "客户信息聚合工作台（ClientHub）产品需求文档"
    props.subject = "多来源客户信息聚合、通话转录、聊天解析、客户 360"
    props.author = "产品团队"
    props.comments = "V1.0 需求基线"

    doc.save(OUT)
    print("docx saved:", OUT)

    # Validate
    check = Document(OUT)
    print("paragraphs:", len(check.paragraphs))
    print("tables:", len(check.tables))
    print("inline_shapes:", len(check.inline_shapes))
    headings = [p.text for p in check.paragraphs if p.style.name.startswith("Heading")]
    print("headings:", len(headings))
    print("first headings:", headings[:8])
    text = "\n".join(p.text for p in check.paragraphs)
    for keyword in ["背景与问题", "功能需求详述", "技术架构", "项目计划", "风险与应对"]:
        print(keyword, "=>", keyword in text)


if __name__ == "__main__":
    main()
